"""Excel workbook output for the Huian research report evidence."""
from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from teaching_console.services.research_report_content import display, number, percent


class ExcelResearchData:
    """Create a readable workbook from existing stored annotations only."""

    header_fill = PatternFill("solid", fgColor="5B9BD5")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(*(Side(style="thin", color="B8C7D9") for _ in range(4)))

    def write(
        self,
        path: Path,
        experiment: Mapping[str, object],
        counts: Sequence[Mapping[str, object]],
        predictions: Sequence[Mapping[str, object]],
        metrics: Mapping[str, object],
        prediction_metrics: Mapping[str, object],
    ) -> None:
        workbook = Workbook()
        info = workbook.active
        info.title = "实验信息"
        self._info_sheet(info, experiment)
        self._count_sheet(workbook.create_sheet("人数标注"), counts)
        self._prediction_sheet(workbook.create_sheet("预测验证"), predictions)
        self._metric_sheet(workbook.create_sheet("统计结果"), metrics, prediction_metrics)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()

    def _info_sheet(self, sheet, experiment: Mapping[str, object]) -> None:
        sheet.append(["项目", "内容"])
        for label, value in [
            ("实验名称", display(experiment.get("name"), "未命名实验")),
            ("实验日期", display(experiment.get("experiment_date") or experiment.get("created_at"))),
            ("参与学生", display(experiment.get("participants"))),
            ("实验目的", display(experiment.get("purpose") or experiment.get("description"))),
            ("测试视频", Path(str(experiment.get("video_path") or "")).name or "尚未填写"),
            ("实验编号", str(experiment.get("id") or "")),
        ]:
            sheet.append([label, value])
        self._finish(sheet, widths=[20, 78])

    def _count_sheet(self, sheet, rows: Sequence[Mapping[str, object]]) -> None:
        sheet.append(["序号", "视频时间（秒）", "帧号", "人工真实人数", "AI检测人数", "相差人数", "观察记录"])
        for row in rows:
            sheet.append([row.get("sample_index"), row.get("video_time_seconds"), row.get("frame_index"), row.get("ground_truth_count"), row.get("system_count"), row.get("absolute_error"), row.get("note", "")])
        self._finish(sheet, widths=[10, 16, 12, 16, 15, 14, 42])

    def _prediction_sheet(self, sheet, rows: Sequence[Mapping[str, object]]) -> None:
        sheet.append(["预测起点（秒）", "预测时长（秒）", "预测人数", "人工真实人数", "相差人数"])
        for row in rows:
            sheet.append([row.get("anchor_time_seconds"), row.get("horizon_seconds"), row.get("prediction_count"), row.get("ground_truth_count"), row.get("absolute_error")])
        self._finish(sheet, widths=[18, 16, 16, 18, 16])

    def _metric_sheet(self, sheet, metrics: Mapping[str, object], prediction_metrics: Mapping[str, object]) -> None:
        sheet.append(["统计项目", "数值", "儿童解释"])
        sheet.append(["标注任务总数", metrics.get("total_tasks", 0), "一共需要完成的人数标注任务。"])
        sheet.append(["已完成人工标注", metrics.get("completed_ground_truth", 0), "已经由同学观察并记录真实人数的任务。"])
        sheet.append(["可评价样本", metrics.get("evaluated_samples", 0), "同时有人工人数和 AI 人数的样本。"])
        sheet.append(["MAE（平均绝对误差）", metrics.get("mae"), "平均每次 AI 和人工相差几个人。"])
        sheet.append(["最大绝对误差", metrics.get("max_absolute_error"), "AI 和人工相差最多的一次。"])
        sheet.append(["完全相同率", percent(metrics.get("exact_match_rate")), "AI 人数与人工人数完全相同的比例。"])
        for horizon in (10, 20, 30):
            values = prediction_metrics if prediction_metrics else {}
            values = values if isinstance(values, Mapping) else {}
            sheet.append([f"+{horizon}秒预测 MAE", values.get(f"mae_{horizon}"), "预测人数与后来人工观察人数的平均相差。"])
        self._finish(sheet, widths=[27, 20, 62])

    def _finish(self, sheet, *, widths: list[int]) -> None:
        for cell in sheet[1]:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = self.border
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

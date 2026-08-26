"""Excel data writer for the child-facing AI difficult-frame experiment."""
from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class AIGrowthExcelData:
    """Write raw difficult-frame observations and their summaries to an XLSX file."""

    header_fill = PatternFill("solid", fgColor="2E74B5")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(*(Side(style="thin", color="B8C7D9") for _ in range(4)))

    def write(
        self,
        path: Path,
        experiment: Mapping[str, object],
        evidence: Sequence[Mapping[str, object]],
        frames: Sequence[Mapping[str, object]],
        verification: str,
    ) -> None:
        workbook = Workbook()
        self._info_sheet(workbook.active, experiment)
        self._frames_sheet(workbook.create_sheet("关键数据记录"), frames)
        self._summary_sheet(workbook.create_sheet("本次实验统计"), evidence, frames)
        self._verification_sheet(workbook.create_sheet("研究结论"), verification)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()

    @staticmethod
    def report_title(experiment: Mapping[str, object]) -> str:
        stem = Path(str(experiment.get("video_path") or "")).stem or str(experiment.get("name") or "未命名视频")
        return f"AI困难帧实验_{stem}"
    def _info_sheet(self, sheet, experiment: Mapping[str, object]) -> None:
        sheet.title = "实验信息"
        sheet.append(["项目", "内容"])
        for label, value in (
            ("报告标题", self.report_title(experiment)),
            ("实验日期", experiment.get("experiment_date") or experiment.get("created_at") or "待填写"),
            ("参与学生", experiment.get("participants") or "待填写"),
            ("测试视频", Path(str(experiment.get("video_path") or "")).name or "待填写"),
            ("研究问题", experiment.get("purpose") or experiment.get("description") or "AI 在哪些画面里容易把人数数错？"),
        ):
            sheet.append([label, value])
        self._finish(sheet, [20, 78])

    def _frames_sheet(self, sheet, frames: Sequence[Mapping[str, object]]) -> None:
        sheet.append(["图片编号", "画面时间（秒）", "AI 原来数", "我框出人数", "相差人数", "完成情况", "我认为为什么难", "AI 推荐关注原因"])
        for row in frames:
            sheet.append([
                row["frame_index"], row["video_time_seconds"], row["system_count"],
                row["human_count"], row["absolute_error"], "已完成框选" if row["annotation_completed"] else "待完成",
                row["student_reason"], row["recommendation_reasons"],
            ])
        self._finish(sheet, [12, 16, 14, 16, 14, 15, 25, 48])

    def _summary_sheet(self, sheet, evidence: Sequence[Mapping[str, object]], frames: Sequence[Mapping[str, object]]) -> None:
        completed = [row for row in frames if row["annotation_completed"]]
        evaluable = [row for row in completed if row["system_count"] is not None and row["human_count"] is not None]
        errors = [float(row["absolute_error"]) for row in evaluable if row["absolute_error"] is not None]
        sheet.append(["统计项目", "数值", "说明"])
        values = (
            ("AI 挑出的画面", sum(int(row["candidate_frames"]) for row in evidence), "AI 从视频中挑出的值得仔细观察的图片。"),
            ("我保留的画面", sum(int(row["kept_frames"]) for row in evidence), "进入人工框选环节的图片。"),
            ("我完成框选", len(completed), "已由孩子保存人工框选答案的图片。"),
            ("我框出的人数", sum(int(row["person_boxes"]) for row in evidence), "孩子在所有已完成图片中框出的人数总和。"),
            ("可比较画面", len(evaluable), "同时具有 AI 原来人数和人工框选人数的图片。"),
            ("平均相差人数", sum(errors) / len(errors) if errors else None, "AI 原来人数与人工框选人数平均相差几个人。"),
            ("最大相差人数", max(errors) if errors else None, "AI 原来人数与人工框选人数相差最多的一次。"),
        )
        for row in values:
            sheet.append(row)
        self._finish(sheet, [24, 18, 66])
        for cell in sheet["B"]:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    def _verification_sheet(self, sheet, verification: str) -> None:
        sheet.append(["AI 成长验证说明"])
        sheet.append([verification])
        self._finish(sheet, [105])
        sheet["A2"].alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[2].height = 75

    def _finish(self, sheet, widths: Sequence[int]) -> None:
        for cell in sheet[1]:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = self.border
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

"""Word report writer for the child-facing AI difficult-frame experiment."""
from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class AIGrowthWordReport:
    """Write a readable report from saved difficult-frame annotations only."""

    def write(
        self,
        path: Path,
        experiment: Mapping[str, object],
        evidence: Sequence[Mapping[str, object]],
        frames: Sequence[Mapping[str, object]],
        reason_text: str,
        verification: str,
    ) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.8)
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(10.5)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(title.add_run(self.report_title(experiment)), 20, bold=True, color="0B2545")

        self._heading(document, "一、我这次想研究什么")
        self._table(document, ["项目", "记录"], [
            ("研究问题", str(experiment.get("purpose") or experiment.get("description") or "AI 在哪些画面里容易把人数数错？")),
            ("研究视频", Path(str(experiment.get("video_path") or "")).name or "待填写"),
            ("实验日期", str(experiment.get("experiment_date") or experiment.get("created_at") or "待填写")),
            ("参与学生", str(experiment.get("participants") or "待填写")),
        ], [1.45, 5.0])

        self._heading(document, "二、我怎样完成这次实验")
        for item in (
            "让 AI 从视频中挑出值得仔细观察的画面。",
            "观察画面，判断 AI 为什么可能数错。",
            "用鼠标逐个框出我看到的人，留下人工答案。",
            "把人工答案和 AI 原来的结果保存在实验数据中。",
        ):
            document.add_paragraph(item, style="List Number")

        self._heading(document, "三、本次实验结果")
        summary_rows = [
            (row["candidate_frames"], row["kept_frames"], row["completed_frames"], row["person_boxes"])
            for row in evidence
        ] or [(0, 0, 0, 0)]
        self._table(document, ["AI 挑出的画面", "我保留的画面", "我完成框选", "我框出的人数"], summary_rows, [1.65, 1.65, 1.65, 1.65])
        document.add_paragraph(f"我观察到 AI 容易出错的原因：{reason_text}")

        self._heading(document, "四、关键数据记录")
        rows = [
            (f"{float(row['video_time_seconds']):.2f}", row["system_count"], row["human_count"],
             row["absolute_error"], row["student_reason"] or "未填写")
            for row in frames
        ] or [("—", "—", "—", "—", "—")]
        self._table(document, ["画面时间（秒）", "AI 原来数", "我框出人数", "相差", "我认为为什么难"], rows, [1.3, 1.3, 1.3, 1.0, 2.4])

        self._heading(document, "五、本次实验结论")
        document.add_paragraph(verification)
        self._heading(document, "六、我的实验心得")
        reflection = str(experiment.get("student_reflection") or "").strip()
        if reflection:
            document.add_paragraph(reflection)
        else:
            document.add_paragraph("我在这次实验中发现：____________________________________________________________")
            document.add_paragraph("我认为 AI 容易数错，是因为：______________________________________________________")
            document.add_paragraph("下一次我想继续研究的问题是：____________________________________________________")

        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)

    @staticmethod
    def report_title(experiment: Mapping[str, object]) -> str:
        stem = Path(str(experiment.get("video_path") or "")).stem or str(experiment.get("name") or "未命名视频")
        return f"AI困难帧实验_{stem}"
    def _heading(self, document: Document, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(5)
        self._font(paragraph.add_run(text), 14, bold=True, color="2E74B5")

    def _table(self, document: Document, headers: Sequence[object], rows: Sequence[Sequence[object]], widths: Sequence[float]) -> None:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False
        for index, value in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.width = Inches(widths[index])
            self._shade(cell, "2E74B5")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._font(cell.paragraphs[0].add_run(str(value)), 8.5, bold=True, color="FFFFFF")
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].width = Inches(widths[index])
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._font(cells[index].paragraphs[0].add_run(str(value if value is not None else "—")), 8.5)

    @staticmethod
    def _font(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    @staticmethod
    def _shade(cell, color: str) -> None:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shading)

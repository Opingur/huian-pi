"""Word rendering for child-friendly Huian research reports."""
from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from teaching_console.services.research_report_content import (
    ReportScreenshot, display, number, percent, research_method,
    research_process_summary, student_response,
)


class WordResearchReport:
    """Create a report from saved evidence, deliberately excluding activity logs."""

    def write(
        self,
        path: Path,
        experiment: Mapping[str, object],
        counts: Sequence[Mapping[str, object]],
        predictions: Sequence[Mapping[str, object]],
        metrics: Mapping[str, object],
        prediction_metrics: Mapping[str, object],
        screenshots: Sequence[ReportScreenshot],
    ) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(1.0)
        section.left_margin = section.right_margin = Inches(1.0)
        section.header_distance = section.footer_distance = Inches(0.492)
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(footer.add_run("慧安楼道｜我的 AI 科研实验报告"), 8, color="666666")

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(5)
        self._font(title.add_run("慧安楼道 AI人数检测实验报告"), 22, bold=True, color="0B2545")
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._font(subtitle.add_run(display(experiment.get("name"), "未命名实验")), 13, color="2E74B5")

        self._heading(document, "一、实验基本信息")
        self._table(document, ["项目", "内容"], [
            ("实验名称", display(experiment.get("name"))),
            ("实验日期", display(experiment.get("experiment_date") or experiment.get("created_at"))),
            ("参与学生", display(experiment.get("participants"))),
            ("测试视频", Path(str(experiment.get("video_path") or "")).name or "尚未填写"),
        ], [1.55, 4.95])

        self._heading(document, "二、实验目的")
        document.add_paragraph(display(experiment.get("purpose") or experiment.get("description")))
        self._heading(document, "三、实验方法")
        for method in research_method():
            document.add_paragraph(method, style="List Number")
        self._heading(document, "四、实验过程")
        document.add_paragraph(research_process_summary(experiment, counts, predictions))

        self._heading(document, "五、实验数据")
        values = [
            (f"{number(row.get('video_time_seconds'), 1)} 秒", display(row.get("ground_truth_count"), "—"),
             display(row.get("system_count"), "—"), number(row.get("absolute_error"), 1))
            for row in counts[:30]
        ] or [("—", "—", "—", "—")]
        self._table(document, ["时间", "人工人数", "AI人数", "误差"], values, [1.6, 1.4, 1.4, 1.2])
        document.add_paragraph(
            f"MAE：{number(metrics.get('mae'), 2)} 人，表示 AI 平均每次和人工结果相差几个人；"
            f"最大误差：{number(metrics.get('max_absolute_error'), 1)} 人；"
            f"完全正确率：{percent(metrics.get('exact_match_rate'))}。"
        )
        self._screenshots(document, screenshots)

        self._student_section(document, "六、我的发现", experiment.get("findings"), "实验中 AI 哪里做得不好？")
        self._student_section(document, "七、为什么会这样？", experiment.get("cause_analysis"), "你认为造成这种结果的原因是什么？")
        self._student_section(document, "八、下一步怎么改进？", experiment.get("next_improvement"), "如果继续研究，你准备怎样帮助 AI 做得更好？")
        if predictions:
            self._heading(document, "九、预测验证")
            values = [
                (f"+{row['horizon_seconds']} 秒", display(row["ground_truth_count"], "—"),
                 number(row["prediction_count"], 1), number(row["absolute_error"], 1))
                for row in predictions[:30]
            ]
            self._table(document, ["预测时间", "真实人数", "预测人数", "误差"], values, [1.55, 1.65, 1.65, 1.65])
            summary = "；".join(
                f"+{horizon} 秒 MAE：{number(prediction_metrics.get(f'mae_{horizon}'), 2)} 人"
                for horizon in (10, 20, 30)
            )
            document.add_paragraph(summary)
        reflection_number = "十" if predictions else "九"
        self._student_section(document, f"{reflection_number}、学生心得", experiment.get("student_reflection"), "这次实验你学到了什么？")
        document.save(path)

    def _screenshots(self, document: Document, screenshots: Sequence[ReportScreenshot]) -> None:
        if not screenshots:
            return
        self._heading(document, "关键截图")
        for screenshot in screenshots:
            try:
                document.add_picture(str(screenshot.path), width=Inches(5.8))
            except OSError:
                continue
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._font(caption.add_run(screenshot.caption), 9, color="555555")

    def _student_section(self, document: Document, heading: str, value: object, prompt: str) -> None:
        self._heading(document, heading)
        document.add_paragraph(student_response(value, prompt))

    @staticmethod
    def _font(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)

    def _heading(self, document: Document, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(6)
        self._font(paragraph.add_run(text), 14, bold=True, color="2E74B5")

    def _table(self, document: Document, headers: Sequence[str], values: Sequence[Sequence[object]], widths: Sequence[float]) -> None:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.width = Inches(widths[index])
            self._shade(cell, "2E74B5")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._font(cell.paragraphs[0].add_run(str(header)), 9, bold=True, color="FFFFFF")
        for row in values:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].width = Inches(widths[index])
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._font(cells[index].paragraphs[0].add_run(str(value)), 9)

    @staticmethod
    def _shade(cell, color: str) -> None:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shading)

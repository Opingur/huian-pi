from __future__ import annotations

import cv2
import numpy as np
from docx import Document
from PIL import Image
from pypdf import PdfReader

import shutil
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from teaching_console.services.research_count_service import ResearchCountService
from teaching_console.services.research_record_exporter import ResearchRecordExporter
from teaching_console.services.research_store import ResearchStore


class ResearchRecordExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.store = ResearchStore(self.root)

    def tearDown(self) -> None:
        shutil.rmtree(self.root)

    def test_legacy_database_is_migrated_and_can_generate_a_report(self) -> None:
        legacy_root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, legacy_root)
        database = legacy_root / "validation" / "research_data" / "huian_research.sqlite3"
        database.parent.mkdir(parents=True)
        connection = sqlite3.connect(database)
        try:
            connection.execute(
                "CREATE TABLE experiments (id TEXT PRIMARY KEY, name TEXT NOT NULL, video_path TEXT NOT NULL, "
                "experiment_type TEXT NOT NULL, description TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, git_commit TEXT)"
            )
            connection.execute(
                "INSERT INTO experiments VALUES (?, ?, ?, ?, ?, ?, ?)",
                ("old", "旧实验", "old.mp4", "teaching", "旧说明", "2026-08-20T10:00:00+00:00", None),
            )
            connection.commit()
        finally:
            connection.close()
        migrated = ResearchStore(legacy_root)
        connection = sqlite3.connect(migrated.database_path)
        try:
            columns = {row[1] for row in connection.execute("PRAGMA table_info(experiments)")}
            tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        finally:
            connection.close()
        self.assertTrue({"purpose", "participants", "experiment_scene", "experiment_date", "findings", "cause_analysis", "next_improvement", "student_reflection"}.issubset(columns))
        self.assertIn("research_activity_log", tables)
        self.assertEqual(migrated.get_experiment("old")["name"], "旧实验")
        exported = ResearchRecordExporter(migrated).export("old", legacy_root / "huian_research_records")
        self.assertTrue(exported.docx_path.is_file())
        self.assertTrue(exported.pdf_path.is_file())
        self.assertTrue(exported.xlsx_path.is_file())

    def test_activity_log_uses_child_facing_research_descriptions(self) -> None:
        experiment_id = self.store.create_experiment("AI人数实验", "stairs.mp4", "teaching")
        annotation_id = self.store.create_count_annotation(experiment_id, 1, 10.0, 150)
        self.store.update_ground_truth(experiment_id, 1, 3, "有人被遮挡")
        self.store.update_system_count(annotation_id, 2)
        self.store.update_experiment_record(experiment_id, findings="人数高峰需要重点观察")
        ResearchCountService(self.store).add_key_sample(experiment_id, 21.0, 15)
        descriptions = [row["description"] for row in self.store.activities(experiment_id)]
        self.assertIn("创建实验：AI人数实验", descriptions)
        self.assertIn("完成第 10.0 秒人数人工标注", descriptions)
        self.assertIn("保存实验记录", descriptions)
        self.assertIn("查看第 10.0 秒的 AI 人数结果", descriptions)
        self.assertIn("添加第 21.0 秒关键样本", descriptions)

    def test_export_creates_word_pdf_excel_raw_evidence_and_embeds_existing_screenshot(self) -> None:
        experiment_id = self.store.create_experiment(
            "AI人数检测实验01", "C:/videos/stairs.mp4", "teaching",
            purpose="比较人工数人与 AI 人数是否一致", participants="小明、小红",
            experiment_scene="学校教学楼二层楼道", experiment_date="2026-08-22",
            findings="遮挡时 AI 容易少数人", cause_analysis="楼道视角下人物互相遮挡",
            next_improvement="补充遮挡场景图片", student_reflection="我学会了先记录真实人数再看 AI 答案。",
        )
        annotation_id = self.store.create_count_annotation(experiment_id, 1, 10.0, 150)
        self.store.update_ground_truth(experiment_id, 1, 3, "两人靠得很近")
        self.store.update_system_count(annotation_id, 2)
        prediction_id = self.store.create_prediction_annotation(experiment_id, 8.0, 120, 2, 0.2, 3.0, 4.0, 5.0)
        self.store.update_prediction_ground_truth(prediction_id, 10, 2)
        self.store.create_model_experiment("AI成长实验01", "huian_person_v1", "models/yolov8n.pt", research_experiment_id=experiment_id)

        root = self.root / "huian_research_records"
        report_dir = ResearchRecordExporter._directory(root, dict(self.store.get_experiment(experiment_id)))
        raw_dir = report_dir / "原始数据"
        raw_dir.mkdir(parents=True)
        raw_dir.joinpath("研究记录摘要.json").write_text(f'{{"experiment_id": "{experiment_id}"}}', encoding="utf-8")
        preview_dir = report_dir / "关键截图"
        preview_dir.mkdir(parents=True)
        Image.new("RGB", (80, 60), (90, 155, 213)).save(preview_dir / "关键样本.png")
        exported = ResearchRecordExporter(self.store).export(experiment_id, root)

        self.assertEqual(exported.directory.name, "AI人数检测实验01_2026-08-22")
        self.assertEqual(exported.data_path.parent.name, "原始数据")
        self.assertTrue(exported.result_path.is_file())
        self.assertTrue(exported.markdown_path.is_file())
        self.assertTrue(exported.docx_path.is_file())
        self.assertTrue(exported.pdf_path.is_file())
        self.assertTrue(exported.xlsx_path.is_file())
        self.assertTrue(exported.screenshots)
        self.assertEqual(exported.pdf_path.read_bytes()[:4], b"%PDF")
        with zipfile.ZipFile(exported.docx_path) as archive:
            self.assertIn("word/document.xml", archive.namelist())
            self.assertTrue(any(name.startswith("word/media/") for name in archive.namelist()))
        workbook = load_workbook(exported.xlsx_path, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["实验信息", "人数标注", "预测验证", "统计结果"])
        self.assertEqual(workbook["人数标注"][2][3].value, 3)
        self.assertEqual(workbook["预测验证"][2][2].value, 3.0)
        workbook.close()
        markdown = exported.markdown_path.read_text(encoding="utf-8")
        self.assertIn("## 三、实验方法", markdown)
        self.assertIn("## 四、实验过程", markdown)
        self.assertIn("## 六、我的发现", markdown)
        self.assertIn("## 九、预测验证", markdown)
        self.assertIn("小明、小红", markdown)
        self.assertIn("遮挡时 AI 容易少数人", markdown)
        self.assertIn("平均每次和人工结果相差几个人", markdown)
        self.assertIn("+10 秒", markdown)
        self.assertNotIn("创建实验：", markdown)
        self.assertNotIn("完成第 10.0 秒人数人工标注", markdown)
        word_text = "\n".join(paragraph.text for paragraph in Document(exported.docx_path).paragraphs)
        self.assertIn("九、预测验证", word_text)
        self.assertNotIn("创建实验：", word_text)
        self.assertNotIn("完成第 10.0 秒人数人工标注", word_text)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(exported.pdf_path).pages)
        self.assertIn("预测验证", pdf_text)
        self.assertNotIn("创建实验", pdf_text)
        summary = exported.result_path.read_text(encoding="utf-8")
        self.assertIn("prediction_records", summary)
        self.assertIn("3,2,1", exported.data_path.read_text(encoding="utf-8-sig"))


    def test_export_generates_named_automatic_screenshots_from_source_video(self) -> None:
        video_path = self.root / "source.avi"
        writer = cv2.VideoWriter(str(video_path), cv2.VideoWriter_fourcc(*"MJPG"), 10.0, (64, 48))
        self.assertTrue(writer.isOpened())
        try:
            for index in range(120):
                writer.write(np.full((48, 64, 3), index % 255, dtype=np.uint8))
        finally:
            writer.release()
        experiment_id = self.store.create_experiment("自动截图实验", video_path, "teaching", experiment_date="2026-08-23")
        annotation_id = self.store.create_count_annotation(experiment_id, 1, 1.0, 10)
        self.store.update_ground_truth(experiment_id, 1, 5, "AI 少数了")
        self.store.update_system_count(annotation_id, 1)
        prediction_id = self.store.create_prediction_annotation(experiment_id, 0.0, 0, 1, 0.1, 1.0, None, None)
        self.store.update_prediction_ground_truth(prediction_id, 10, 7)
        exported = ResearchRecordExporter(self.store).export(experiment_id, self.root / "records")
        names = {path.name for path in (exported.directory / "关键截图").iterdir()}
        self.assertTrue({"001_最大误差.png", "002_AI漏检.png", "003_预测误差最大.png"}.issubset(names))
        self.assertTrue(all(item.path.is_file() for item in exported.screenshots))

if __name__ == "__main__":
    unittest.main()

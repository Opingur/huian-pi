from __future__ import annotations

import shutil
import tempfile
import unittest
from docx import Document
from openpyxl import load_workbook
from pathlib import Path

from teaching_console.services.ai_growth_record_exporter import AIGrowthRecordExporter
from teaching_console.services.research_store import ResearchStore


class AIGrowthRecordExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.root)
        self.store = ResearchStore(self.root)

    def test_legacy_split_name_is_exported_as_child_friendly_difficult_frame_video(self) -> None:
        experiment_id = self.store.create_experiment("旧分组报告", "legacy.mp4", "teaching")
        project_id = self.store.create_detection_annotation_project("旧检查视频", "legacy.mp4", "val", "huian_person_v1")
        exported = AIGrowthRecordExporter(self.store).export(experiment_id, [project_id], output_root=self.root / "records")

        markdown = exported.markdown_path.read_text(encoding="utf-8")
        summary = __import__("json").loads(exported.summary_path.read_text(encoding="utf-8"))

        self.assertIn("困难帧实验视频", markdown)
        self.assertNotIn("老师检查图片", markdown)
        self.assertEqual(summary["projects"][0]["purpose"], "困难帧实验视频")
    def test_markdown_export_uses_saved_boxes_and_real_ab_result(self) -> None:
        experiment_id = self.store.create_experiment("AI成长测试", "learn.mp4", "teaching", purpose="帮助 AI 数清楼道人数")
        project_id = self.store.create_detection_annotation_project("练习", "learn.mp4", "train", "huian_person_v1")
        frame_id = self.store.create_detection_frame_annotation(project_id, 1, 0.1, 640, 480, student_reason="有人被遮挡", annotation_completed=True)
        self.store.create_detection_person_box(frame_id, 1, 1, 20, 30)
        exported = AIGrowthRecordExporter(self.store).export(
            experiment_id, [project_id], output_root=self.root / "records",
            ab_result={"baseline": {"count_mae": 3.0}, "candidate": {"count_mae": 1.0}},
        )
        content = exported.markdown_path.read_text(encoding="utf-8")
        self.assertIn("有人被遮挡：1 张", content)
        self.assertIn("AI 进步了", content)
        self.assertTrue(exported.summary_path.is_file())
        self.assertTrue(exported.docx_path.is_file())
        self.assertTrue(exported.xlsx_path.is_file())
        document_text = "\n".join(paragraph.text for paragraph in Document(exported.docx_path).paragraphs)
        self.assertIn("AI困难帧实验_learn", document_text)
        self.assertNotIn("慧安楼道 AI 成长实验记录", document_text)
        self.assertIn("我的实验心得", document_text)

        workbook = load_workbook(exported.xlsx_path, data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["实验信息", "关键数据记录", "本次实验统计", "研究结论"])
        self.assertEqual(workbook["关键数据记录"]["A2"].value, 1)


if __name__ == "__main__":
    unittest.main()

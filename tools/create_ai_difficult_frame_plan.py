"""Generate the student-fillable "AI 困难帧实验方案" Word form.

This is a blank planning sheet for the AI growth / difficult-frame experiment.
It only asks the child to plan and guess BEFORE the experiment.  It never writes
a conclusion, a cause answer, or any fabricated metric — those must come from the
real experiment afterwards.

Reuses the project's established Word conventions (Microsoft YaHei, A4, the
Huian orange #FF9423) so it matches the existing research reports.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft YaHei"
ORANGE = "FF9423"   # 慧安橙，来自 rpi_app/ui/dashboard_layout.py 的 _ORANGE
INK = "222222"
MUTED = "666666"
LINE = "B0B0B0"

OUTPUT_NAME = "AI困难帧实验方案_学生填写版.docx"


def _font(run, size: float, *, bold: bool = False, color: str = INK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _spacing(paragraph, before: float = 0, after: float = 0, line: float | None = None) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if line is not None:
        paragraph.paragraph_format.line_spacing = line


def _fill_line(document: Document, after: float = 2) -> None:
    """A single handwriting line (paragraph with a light bottom border)."""
    p = document.add_paragraph()
    _spacing(p, after=after, line=1.35)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LINE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _heading(document: Document, text: str) -> None:
    p = document.add_paragraph()
    _spacing(p, before=9, after=3)
    _font(p.add_run(text), 12.5, bold=True, color=ORANGE)


def _hint(document: Document, text: str, *, after: float = 3) -> None:
    p = document.add_paragraph()
    _spacing(p, after=after, line=1.15)
    _font(p.add_run(text), 9, color=MUTED)


def _field_label(document: Document, label: str) -> None:
    p = document.add_paragraph()
    _spacing(p, after=1, line=1.35)
    _font(p.add_run(label), 10.5, color=INK)


def _shade(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _data_table(document: Document) -> None:
    headers = ["画面/样本编号", "人工真实人数", "AI检测人数", "相差人数", "是否为困难帧", "我的备注"]
    widths = [3.0, 2.4, 2.2, 2.0, 2.4, 5.4]  # cm, sums to 17.4 = usable width
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        _shade(cell, "F5F5F5")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(cell.paragraphs[0].add_run(header), 9, bold=True, color=INK)
    for _ in range(3):  # 3 blank rows to fill by hand later
        cells = table.add_row().cells
        for index, width in enumerate(widths):
            cells[index].width = Cm(width)
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(cells[index].paragraphs[0], after=6, line=1.3)


def _checkbox(document: Document, text: str) -> None:
    p = document.add_paragraph()
    _spacing(p, after=2, line=1.3)
    _font(p.add_run("□ "), 11, color=INK)
    _font(p.add_run(text), 10.5, color=INK)


def build(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)

    # 标题
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(title, after=1)
    _font(title.add_run("AI困难帧实验方案"), 17, bold=True, color=INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(subtitle, after=4)
    _font(subtitle.add_run("慧安楼道｜AI成长实验"), 9, color=MUTED)

    # 基本信息
    _heading(document, "基本信息")
    _field_label(document, "姓名：______________________________    日期：____________________")
    _field_label(document, "实验名称：____________________________________________________")
    _hint(document, "例如：AI在哪些画面中容易数错人数？")

    # 一、我想研究的问题
    _heading(document, "一、我想研究的问题")
    _hint(document, "这次实验，你最想弄清楚AI的什么问题？")
    for _ in range(2):
        _fill_line(document)

    # 二、我的猜想
    _heading(document, "二、我的猜想")
    _hint(document, "实验开始前，先猜一猜：什么情况下AI可能更容易数错人？为什么？")
    for _ in range(3):
        _fill_line(document)
    _hint(document, "可以从遮挡、人物大小、画面位置等角度思考，但没有标准答案。", after=1)

    # 三、我准备怎么做
    _heading(document, "三、我准备怎么做")
    for step in (
        "选择一段测试视频",
        "找出一些具有代表性的画面",
        "先不看AI答案，人工数出真实人数",
        "再查看AI检测人数",
        "找出AI数错的困难画面",
        "记录误差，并分析可能原因",
    ):
        p = document.add_paragraph()
        _spacing(p, after=1, line=1.25)
        _font(p.add_run(step), 10.5, color=INK)
    p = document.add_paragraph()
    _spacing(p, after=2, line=1.3)
    _font(p.add_run("我还准备："), 10.5, color=INK)
    _font(p.add_run("____________________________________"), 10.5, color=INK)

    # 我准备记录哪些数据
    _heading(document, "我准备记录哪些数据")
    _data_table(document)
    _hint(document, "（先留空，实验时再填写真实数据。）", after=2)

    # 四、我重点想观察什么
    _heading(document, "四、我重点想观察什么")
    for item in ("人互相遮挡时", "人物比较小时", "人在画面边缘时", "多个人靠得很近时"):
        _checkbox(document, item)
    _checkbox(document, "其他：____________________________")
    _hint(document, "上面只是我准备观察的方向，不代表AI一定会出错。", after=2)

    # 实验完成后我要回答的问题
    _heading(document, "实验完成后我要回答的问题")
    for question in (
        "1. AI在哪些画面里表现得比较好？",
        "2. AI在哪些画面里容易出现误差？",
        "3. 我认为可能是什么原因？",
        "4. 如果继续改进，我想先改什么？",
    ):
        p = document.add_paragraph()
        _spacing(p, after=1, line=1.25)
        _font(p.add_run(question), 10.5, color=INK)

    document.save(path)
    return path


def main() -> None:
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent
    output_dir = project_root / "慧安楼道安全监测系统_研究资料" / "08_学生活动记录"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = build(output_dir / OUTPUT_NAME)

    # 最小自检：文件能重新打开、结构完整。
    reopened = Document(str(path))
    assert any("AI困难帧实验方案" in p.text for p in reopened.paragraphs), "缺少标题"
    assert len(reopened.tables) == 1, f"应有 1 张数据表，实际 {len(reopened.tables)}"
    print(f"已生成：{path}")
    print(f"段落数：{len(reopened.paragraphs)}，表格数：{len(reopened.tables)}")


if __name__ == "__main__":
    main()
